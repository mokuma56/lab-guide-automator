"""Build the interactive AI Canvas Prompt Library app from the prompt-testing spreadsheet.

    uv run --with openpyxl python3 prompt_library/build_prompt_library.py \
        "<path to dCloud Canvas Prompt Library.xlsx>" --guide <guide id> [--ratings-url <apps script url>]

Writes data/apps/<guide id>/prompt-library.html, which the MkDocs export copies to docs/apps/.
Only public-safe fields are included: run outputs, trace IDs, AI consistency commentary, and
internal source names stay out of the published app.
"""
from __future__ import annotations

import argparse
import json
import re
import statistics
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
ROOT = HERE.parent


def _seconds(value) -> int | None:
    m = re.match(r"\s*(\d+)m\s*(\d+)s", str(value or ""))
    return int(m[1]) * 60 + int(m[2]) if m else None


def _problem_ids(wb) -> set[str]:
    """Prompts the Summary sheet lists as PROBLEMATIC / UNRESOLVED (repeated timeouts or failures)."""
    ids = set()
    for row in wb["Summary"].iter_rows(values_only=True):
        if row[0] and re.fullmatch(r"P\d{3}", str(row[0])) and "PROBLEMATIC" in str(row[2] or ""):
            ids.add(row[0])
    return ids


def load_prompts(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    problems = _problem_ids(wb)
    rows = list(wb["Prompts"].iter_rows(values_only=True))
    header = rows[0]
    prompts = []
    for raw in rows[1:]:
        r = dict(zip(header, raw))
        if not r["ID"]:
            continue
        grades = [int(g) for i in range(1, 6) if str(g := r[f"Run {i} Grade (1-5)"]).strip().isdigit()]
        durations = [d for i in range(1, 6) if (d := _seconds(r[f"Run {i} Duration"]))]
        domain = str(r["Domain"] or "Other").strip()
        consistency = r["Consistency Rating (AI: High/Medium/Low)"]
        prompts.append({
            "id": r["ID"],
            "parent": r["Follow-up To (Parent ID)"] or None,
            "trigger": str(r["Follow-up Trigger / Condition"]).strip() if r["Follow-up Trigger / Condition"] else None,
            "domain": domain,
            "category": domain.split("/")[0].strip(),
            "title": str(r["Business Framing"] or "").strip(),
            "prompt": str(r["Prompt Text"] or "").strip(),
            "expect": str(r["Expected Signal / What It Should Correlate"] or "").strip(),
            "consistency": str(consistency).upper() if consistency else None,
            "grade": round(statistics.mean(grades), 2) if grades else None,
            "graded": len(grades),
            "secs": round(statistics.mean(durations)) if durations else None,
            "runs": sum(1 for i in range(1, 6) if r[f"Run {i} Output"]),
            "problem": r["ID"] in problems or "PROBLEMATIC" in str(r["Notes"] or ""),
        })
    return prompts


def build(xlsx: Path, guide_id: str, ratings_url: str = "") -> Path:
    data = json.dumps(load_prompts(xlsx), ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = (HERE / "app.template.html").read_text()
    html = html.replace("__DATA__", data).replace("__RATINGS_URL__", json.dumps(ratings_url))
    out = ROOT / "data" / "apps" / guide_id / "prompt-library.html"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(html)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--guide", required=True, help="Guide ID the app is embedded in")
    ap.add_argument("--ratings-url", default="", help="Google Apps Script web app URL for shared ratings")
    args = ap.parse_args()
    path = build(args.xlsx, args.guide, args.ratings_url)
    print(f"Wrote {path} ({path.stat().st_size // 1024} KB)")
